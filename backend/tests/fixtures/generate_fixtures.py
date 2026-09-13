"""生成测试夹具。

测试夹具是提交进版本库的（体积很小），这个脚本用于在需要时重新生成。
运行：``uv run python tests/fixtures/generate_fixtures.py``

三个夹具分别覆盖三种关键场景：

============================  ==================================================
contract.pdf                  有文本层的中文 PDF（正常路径）
resume.docx                   含表格的中文 Word（验证表格转 Markdown 的逻辑）
scanned.pdf                   没有文本层的 PDF（验证扫描件识别与友好提示）
============================  ==================================================

``contract.pdf`` 依赖系统的 ``cupsfilter``（macOS / 带 CUPS 的 Linux 自带）把
``contract.txt`` 转成嵌入了中文字体的 PDF；找不到该命令时会跳过并保留已有文件。
"""

from __future__ import annotations

import shutil
import subprocess
import sys
from pathlib import Path

FIXTURES_DIR = Path(__file__).parent

CONTRACT_TEXT = """采购合同

甲方：北京星辰科技有限公司
乙方：上海云图信息技术有限公司

第一条 合同标的
乙方向甲方提供企业级数据中台建设服务，包括数据采集、清洗、建模与可视化。

第二条 合同金额
本合同总金额为人民币 1,280,000 元（大写：壹佰贰拾捌万元整）。

第三条 付款方式
合同签订后 5 个工作日内，甲方支付合同总额的 30% 作为预付款。

第四条 交付时间
乙方应于 2026 年 6 月 30 日前完成全部交付内容。

甲方（盖章）：北京星辰科技有限公司
法定代表人：张伟
签订日期：2026年3月15日

乙方（盖章）：上海云图信息技术有限公司
法定代表人：李娜
签订日期：2026年3月15日
"""

RESUME_ROWS = [
    ("时间", "公司", "职位", "主要工作"),
    (
        "2021.03 - 至今",
        "上海云图信息技术有限公司",
        "高级前端工程师",
        "主导企业级数据中台前端架构设计，支撑 12 个业务系统接入；"
        "推动构建体系升级，打包耗时从 8 分钟降至 90 秒。",
    ),
    (
        "2017.07 - 2021.02",
        "北京星辰科技有限公司",
        "前端工程师",
        "负责电商中台订单模块开发，日均处理订单 30 万笔；"
        "建立前端监控体系，线上问题平均定位时间从 2 小时缩短到 15 分钟。",
    ),
    (
        "2015.09 - 2017.06",
        "杭州某网络科技有限公司",
        "初级前端工程师",
        "参与企业官网与后台管理系统开发。",
    ),
]


def _write_contract_text() -> Path:
    path = FIXTURES_DIR / "contract.txt"
    path.write_text(CONTRACT_TEXT, encoding="utf-8")
    return path


def _generate_contract_pdf() -> None:
    """用 cupsfilter 把中文文本渲染成带嵌入字体的 PDF。"""
    cupsfilter = shutil.which("cupsfilter")
    if not cupsfilter:
        print("  跳过 contract.pdf：未找到 cupsfilter 命令", file=sys.stderr)
        return

    source = _write_contract_text()
    target = FIXTURES_DIR / "contract.pdf"
    with target.open("wb") as out:
        subprocess.run(
            [cupsfilter, "-t", "text/plain", str(source)],
            stdout=out,
            stderr=subprocess.DEVNULL,
            check=True,
        )
    print(f"  已生成 {target.name} ({target.stat().st_size} 字节)")


def _generate_resume_docx() -> None:
    """含段落与表格的简历。表格用于验证 Markdown 还原逻辑。"""
    from docx import Document

    document = Document()

    document.add_heading("个人简历", level=1)

    document.add_heading("基本信息", level=2)
    document.add_paragraph("姓名：王小明")
    document.add_paragraph("求职意向：AI 全栈应用开发工程师")
    document.add_paragraph("联系电话：138-0000-0000")
    document.add_paragraph("邮箱：wangxiaoming@example.com")

    document.add_heading("工作经历", level=2)
    table = document.add_table(rows=len(RESUME_ROWS), cols=4)
    table.style = "Table Grid"
    for row_index, row_data in enumerate(RESUME_ROWS):
        for col_index, value in enumerate(row_data):
            table.cell(row_index, col_index).text = value

    document.add_heading("技能清单", level=2)
    document.add_paragraph("前端：Vue 3、TypeScript、Vite、Pinia、Ant Design Vue")
    document.add_paragraph("后端：Python、FastAPI、PostgreSQL、SQLAlchemy、Docker")
    document.add_paragraph("AI：DeepSeek API、Prompt 工程、RAG、Function Calling")

    document.add_heading("教育背景", level=2)
    document.add_paragraph("2011.09 - 2015.06 某某大学 计算机科学与技术 本科")

    path = FIXTURES_DIR / "resume.docx"
    document.save(str(path))
    print(f"  已生成 {path.name} ({path.stat().st_size} 字节)")


def _generate_scanned_pdf() -> None:
    """一份只有空白页、没有文本层的 PDF，模拟扫描件。"""
    import pypdfium2 as pdfium

    path = FIXTURES_DIR / "scanned.pdf"
    document = pdfium.PdfDocument.new()
    try:
        document.new_page(595, 842)  # A4
        with path.open("wb") as out:
            document.save(out)
    finally:
        document.close()
    print(f"  已生成 {path.name} ({path.stat().st_size} 字节)")


def main() -> None:
    print("生成测试夹具到", FIXTURES_DIR)
    _generate_contract_pdf()
    _generate_resume_docx()
    _generate_scanned_pdf()
    _write_contract_text()
    print("完成")


if __name__ == "__main__":
    main()
