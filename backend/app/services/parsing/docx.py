"""Word (.docx) 文本提取。

关键点：**表格必须按 Markdown 表格还原**。

python-docx 的 ``document.paragraphs`` 只会返回正文段落，**完全跳过表格内容**。
发票、简历这类文档的价值恰恰大量存在于表格里——直接丢掉表格，
模型就只能靠猜，抽取质量会断崖式下跌。

因此这里按文档流顺序遍历 body 的子元素，段落与表格混排还原，
表格渲染成 Markdown 管道表格：模型对 Markdown 表格的解析能力远好于
一串制表符或空格分隔的裸文本。
"""

from __future__ import annotations

import io

from app.core.errors import CorruptedFileError
from app.services.parsing.normalize import normalize_text


def _cell_text(cell) -> str:
    """单元格文本。单元格内可能有多段，用空格连接并转义竖线。"""
    parts = [paragraph.text.strip() for paragraph in cell.paragraphs]
    text = " ".join(part for part in parts if part)
    # 竖线是 Markdown 表格的列分隔符，出现在内容里会破坏表结构
    return text.replace("|", "\\|")


def _table_to_markdown(table) -> str:
    """把 Word 表格渲染成 Markdown 管道表格。"""
    rows = [[_cell_text(cell) for cell in row.cells] for row in table.rows]
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        return ""

    column_count = max(len(row) for row in rows)

    def pad(row: list[str]) -> list[str]:
        return row + [""] * (column_count - len(row))

    lines = ["| " + " | ".join(pad(rows[0])) + " |"]
    lines.append("| " + " | ".join(["---"] * column_count) + " |")
    lines.extend("| " + " | ".join(pad(row)) + " |" for row in rows[1:])
    return "\n".join(lines)


def _iter_block_items(document):
    """按文档流顺序产出 Paragraph 与 Table。

    python-docx 没有提供「按顺序遍历正文」的公开 API，
    只能下探到 XML 层按 body 的子元素标签判断类型。
    这是官方 issue 里长期推荐的写法。
    """
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def extract_docx(data: bytes) -> tuple[str, int | None, str]:
    """提取 .docx 文本。

    :returns: ``(清洗后的文本, 页数, 解析器名)``。
              .docx 没有固定分页的概念（分页由渲染器决定），因此页数返回 None。
    :raises CorruptedFileError: 不是合法的 docx（例如把 .doc 改名成 .docx）
    """
    try:
        from docx import Document as DocxDocument
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover
        raise CorruptedFileError("python-docx 未安装，无法解析 .docx") from exc

    try:
        document = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise CorruptedFileError(
            "Word 文档解析失败。注意：旧版 .doc 格式不受支持，请另存为 .docx 后重试。",
            details={"reason": f"{type(exc).__name__}: {exc}"},
        ) from exc

    blocks: list[str] = []
    for block in _iter_block_items(document):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                blocks.append(text)
        elif isinstance(block, Table):
            markdown = _table_to_markdown(block)
            if markdown:
                blocks.append(markdown)

    return normalize_text("\n\n".join(blocks)), None, "python-docx"
